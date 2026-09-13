import io
import logging
from typing import Dict, Any, List

logger = logging.getLogger(__name__)

def extract_text_from_file(file_bytes: bytes, filename: str, mime_type: str = "") -> Dict[str, Any]:
    """
    Extracts text content and page-level metadata from PDF, DOCX, PPTX, or TXT files.
    """
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    pages = []
    full_text = ""

    try:
        if ext == "pdf" or "pdf" in mime_type:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages, 1):
                txt = page.extract_text() or ""
                pages.append({"page_number": i, "text": txt.strip()})
                full_text += f"\n--- Page {i} ---\n" + txt

        elif ext == "docx" or "word" in mime_type:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            pages.append({"page_number": 1, "text": full_text})

        elif ext == "pptx" or "powerpoint" in mime_type:
            import pptx
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(prs.slides, 1):
                slide_txt = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_txt.append(shape.text)
                txt = "\n".join(slide_txt)
                pages.append({"page_number": i, "text": txt.strip()})
                full_text += f"\n--- Slide {i} ---\n" + txt

        else: # TXT or Fallback
            full_text = file_bytes.decode("utf-8", errors="ignore")
            pages.append({"page_number": 1, "text": full_text})

    except Exception as e:
        logger.error(f"Error parsing file '{filename}': {e}")
        full_text = file_bytes.decode("utf-8", errors="ignore")
        pages.append({"page_number": 1, "text": full_text})

    return {
        "full_text": full_text.strip(),
        "pages": pages,
        "total_pages": len(pages)
    }
